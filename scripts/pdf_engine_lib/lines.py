"""
Detecting a form's own ruled lines (as opposed to a Type3 glyph's outline or
a real table's borders) and rendering/placing them as a background image
behind pdf2docx's reconstructed text - see `convert_to_docx`'s dense-vector-
lines branch in `docx.py` for why this exists at all.
"""
from .util import io_bytes


def _is_ruling_line(path: dict) -> bool:
    """
    True if `path` (one entry from `page.get_drawings()`) is a genuine
    ruling - an underline, a signature line, a table rule - rather than a
    glyph's own outline.

    This distinction only matters at all because of Type3 fonts: a normal
    font's glyphs are drawn with text operators and never show up in
    `get_drawings()` in the first place, but a Type3 glyph IS a tiny vector
    drawing (that is what "Type3" means - each glyph is its own draw
    program) - so on a Type3-font page, "every vector drawing" includes
    every letter of every word. Confirmed directly on this exact class of
    PDF: naively rendering all drawings duplicated the entire page's body
    text, in the original font, floating behind the real reconstructed
    text.

    Two checks tell a ruling apart from a glyph, together, because neither
    alone works: a ruling is THIN in one dimension (a hairline stroke or a
    thin filled bar) which a page number's digit is not (it's small in
    both), and a ruling is also LONG in the other dimension, which a small
    dot-leader period is not (it's small in both dimensions too, just like
    a digit) - so a bare "is thin" check alone would keep leader dots, and
    a bare "is long" check alone would keep multi-character text runs.
    `item_count` is a belt-and-suspenders cap: every ruling observed here
    is drawn with a handful of path segments, where a Type3 glyph cluster -
    especially pdf2docx/PyMuPDF's own merged runs of several adjacent
    glyphs - routinely reaches into the hundreds.
    """
    rect = path['rect']
    width, height = rect.width, rect.height
    return len(path['items']) <= 30 and min(width, height) <= 3.0 and max(width, height) >= 12.0


def _pdf_has_dense_vector_lines(input_path: str) -> bool:
    """
    True if some page draws a lot of ruling lines that are overwhelmingly
    ONE orientation - the specific shape of a form's underline rules
    (`Tələbə: _______`, one per field, all horizontal, nothing dividing
    them into columns), and the case pdf2docx's own table-detection
    heuristic (grouping intersecting borders into cells) gets wrong: it
    picks up those rules as table borders anyway, misgroups them with
    unrelated text runs, and recolours/duplicates spans in the process
    (observed directly on an Azerbaijani internship-report template - text
    that also has no underline in the source coming out red and
    underlined).

    A real table - checked directly against a PDF built entirely out of
    them - draws vertical rulings at several DISTINCT x-positions (one per
    column boundary; every real table sampled had at least 6), and
    pdf2docx's table detection handles that shape correctly; disabling it
    there would trade a working table for the same plain-paragraphs-plus-
    line-image fallback this function exists to avoid needing in the first
    place. Distinct positions rather than a raw vertical count: a resume
    with one single vertical divider line (separating a date column from
    the rest, say) running down the whole page draws that ONE line as
    several separate drawing objects, one per row it passes - plenty of
    "vertical rulings" by count, but still only one real column boundary,
    and pdf2docx's table detection turns that into a bordered two-column
    table that does not exist in the source at all (confirmed directly:
    a resume whose PDF has no table borders anywhere came out with visible
    cell borders and text spilling outside them in Word). So this checks
    for the LOPSIDED case specifically - plenty of rulings, but either
    essentially all one orientation, or with fewer than 2 distinct column
    boundaries - rather than just "plenty of rulings": a raw count alone
    flagged a real, working table too, and disabled it for no reason.

    Filtered through `_is_ruling_line` rather than counting every
    `get_drawings()` entry: on a Type3-font page (see `_is_ruling_line`),
    the raw count is dominated by glyph paths and would trigger this on
    every Type3 PDF regardless of whether it draws any real ruling at all.

    Checked as the BUSIEST single page, not a document-wide average: a
    real-world report is exactly the shape "one line-heavy cover/signature
    page plus several pages of plain prose with zero rulings" - averaging
    over every page (tried first, against a real 10-page report whose
    cover page alone draws 10 rulings) dilutes that one page below any
    sane threshold and never triggers at all. The overlay this gates is
    applied per page regardless, so nothing but the trigger itself needed
    to change.
    """
    import fitz  # PyMuPDF

    document = fitz.open(input_path)
    try:
        if document.page_count == 0:
            return False
        for page in document:
            lines = [path for path in page.get_drawings() if _is_ruling_line(path)]
            if len(lines) < 6:
                continue
            vertical_x_positions = {
                round(path['rect'].x0) for path in lines if path['rect'].height > path['rect'].width
            }
            if len(vertical_x_positions) < 2:
                return True
        return False
    finally:
        document.close()


def _render_lines_only_png(page, matrix) -> bytes:
    """
    Rasterise only `page`'s vector graphics (the ruled lines/boxes a form
    draws for its own "table") onto a transparent background, with no text -
    the twin image `_add_line_overlay_backgrounds` drops behind pdf2docx's
    reconstructed text so a line-heavy PDF keeps its exact ruling without
    pdf2docx ever having to (mis)interpret those lines as table borders.

    Redrawn onto a fresh, contentless page rather than clipping the
    original: a page with nothing painted on it renders as fully
    transparent (`alpha=True`), which a clipped screenshot of the real page
    is not - that would carry the white page background and the text along
    with it.
    """
    import fitz  # PyMuPDF

    overlay_doc = fitz.open()
    try:
        overlay_page = overlay_doc.new_page(width=page.rect.width, height=page.rect.height)
        shape = overlay_page.new_shape()
        for path in page.get_drawings():
            if not _is_ruling_line(path):
                continue
            for item in path['items']:
                op = item[0]
                if op == 'l':
                    shape.draw_line(item[1], item[2])
                elif op == 're':
                    shape.draw_rect(item[1])
                elif op == 'qu':
                    shape.draw_quad(item[1])
                elif op == 'c':
                    shape.draw_bezier(item[1], item[2], item[3], item[4])
            shape.finish(
                color=path.get('color'),
                fill=path.get('fill'),
                width=path.get('width') or 1.0,
                closePath=path.get('closePath', True),
            )
        shape.commit()
        pixmap = overlay_page.get_pixmap(matrix=matrix, alpha=True)
        return pixmap.tobytes('png')
    finally:
        overlay_doc.close()


def _float_picture_behind_text(run, width_emu: int, height_emu: int) -> None:
    """
    Rewrite the inline picture `run.add_picture` just created (its usual,
    only mode) into a floating one anchored to the page at (0, 0) and sent
    behind the text - turning it from "an image inline in the paragraph
    text" into "a page background" without leaving flow at all, so it never
    pushes the real, editable text pdf2docx produced out of place.

    Doing this by mutating the `wp:inline` element `add_picture` already
    built - reusing its `a:graphic`/`pic:pic`/relationship rather than
    constructing a picture from scratch - is deliberate: that element is
    already a complete, valid picture; only the wrapper that positions it
    within the run needs to change from `wp:inline` (flows with text) to
    `wp:anchor` (positioned independently, `behindDoc="1"`).
    """
    from docx.oxml.ns import qn

    drawing = run._r.find(qn('w:drawing'))
    inline = drawing.find(qn('wp:inline'))
    graphic = inline.find(qn('a:graphic'))
    doc_pr = inline.find(qn('wp:docPr'))
    inline.remove(graphic)
    inline.remove(doc_pr)
    drawing.remove(inline)

    from docx.oxml import OxmlElement

    anchor = OxmlElement('wp:anchor')
    anchor.set('behindDoc', '1')
    anchor.set('locked', '0')
    anchor.set('layoutInCell', '0')
    anchor.set('allowOverlap', '1')
    anchor.set('relativeHeight', '0')
    anchor.set('simplePos', '0')

    simple_pos = OxmlElement('wp:simplePos')
    simple_pos.set('x', '0')
    simple_pos.set('y', '0')
    anchor.append(simple_pos)

    position_h = OxmlElement('wp:positionH')
    position_h.set('relativeFrom', 'page')
    offset_h = OxmlElement('wp:posOffset')
    offset_h.text = '0'
    position_h.append(offset_h)
    anchor.append(position_h)

    position_v = OxmlElement('wp:positionV')
    position_v.set('relativeFrom', 'page')
    offset_v = OxmlElement('wp:posOffset')
    offset_v.text = '0'
    position_v.append(offset_v)
    anchor.append(position_v)

    extent = OxmlElement('wp:extent')
    extent.set('cx', str(width_emu))
    extent.set('cy', str(height_emu))
    anchor.append(extent)

    effect_extent = OxmlElement('wp:effectExtent')
    for side in ('l', 't', 'r', 'b'):
        effect_extent.set(side, '0')
    anchor.append(effect_extent)

    anchor.append(OxmlElement('wp:wrapNone'))
    anchor.append(doc_pr)
    anchor.append(graphic)

    drawing.append(anchor)


def _paragraph_ends_section(paragraph_element) -> bool:
    """
    True if this `w:p` carries a section break - Word nests `w:sectPr`
    inside a paragraph's `w:pPr` to mark it as the LAST paragraph of a
    section, everywhere except the document's final section, whose
    `sectPr` sits on the body itself rather than on any paragraph.
    """
    from docx.oxml.ns import qn

    p_pr = paragraph_element.find(qn('w:pPr'))
    return p_pr is not None and p_pr.find(qn('w:sectPr')) is not None


def _add_line_overlay_backgrounds(docx_path: str, input_path: str) -> None:
    """
    Drop each page's `_render_lines_only_png` behind the text pdf2docx
    already reconstructed - see `convert_to_docx`'s dense-vector-lines
    branch for why the text was reconstructed with table detection off in
    the first place, and this module's docstring-level notes above for the
    corruption that combination avoids.

    Assumes pdf2docx keeps its usual one-section-per-source-page pagination
    (true for the single-column reports this branch targets) and maps
    overlay images to sections by that order; a source with multi-column
    sections that split unevenly would misalign past the first mismatch,
    which is why this is gated behind `_pdf_has_dense_vector_lines` rather
    than applied unconditionally.
    """
    import fitz  # PyMuPDF
    from docx import Document
    from docx.shared import Emu

    dpi = 150
    zoom = dpi / 72.0
    matrix = fitz.Matrix(zoom, zoom)
    emu_per_point = 914400 / 72.0

    document = fitz.open(input_path)
    try:
        overlays = [_render_lines_only_png(page, matrix) for page in document]
        page_sizes = [(page.rect.width, page.rect.height) for page in document]
    finally:
        document.close()

    output = Document(docx_path)
    paragraphs = output.paragraphs

    section_starts = [0]
    for index, paragraph in enumerate(paragraphs):
        if _paragraph_ends_section(paragraph._p) and index + 1 < len(paragraphs):
            section_starts.append(index + 1)

    for page_index, start in enumerate(section_starts):
        if page_index >= len(overlays) or start >= len(paragraphs):
            break
        width_pt, height_pt = page_sizes[page_index]
        width_emu = Emu(int(round(width_pt * emu_per_point)))
        height_emu = Emu(int(round(height_pt * emu_per_point)))

        run = paragraphs[start].add_run()
        run.add_picture(io_bytes(overlays[page_index]), width=width_emu, height=height_emu)
        _float_picture_behind_text(run, int(width_emu), int(height_emu))

    output.save(docx_path)
